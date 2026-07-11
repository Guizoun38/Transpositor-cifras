import io

from docx import Document

from application.conversion import convert_musical_text
from application.ports import DocumentProcessor


def _iter_paragraphs(document):
    def from_table(table):
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from from_table(nested)
    yield from document.paragraphs
    for table in document.tables:
        yield from from_table(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from from_table(table)


def _replace_runs(paragraph, new_text):
    old_text = "".join(run.text for run in paragraph.runs)
    if old_text == new_text or not paragraph.runs:
        return
    # Prefix/suffix trimming confines edits to the smallest changed span.
    prefix = 0
    while prefix < min(len(old_text), len(new_text)) and old_text[prefix] == new_text[prefix]:
        prefix += 1
    suffix = 0
    while suffix < min(len(old_text) - prefix, len(new_text) - prefix) and old_text[-1-suffix] == new_text[-1-suffix]:
        suffix += 1
    offsets, pos = [], 0
    for run in paragraph.runs:
        offsets.append(pos); pos += len(run.text)
    start, end = prefix, len(old_text) - suffix
    affected = [i for i, off in enumerate(offsets) if off + len(paragraph.runs[i].text) > start and off < end]
    if not affected:
        return
    first, last = affected[0], affected[-1]
    before = paragraph.runs[first].text[:start-offsets[first]]
    after = paragraph.runs[last].text[end-offsets[last]:]
    paragraph.runs[first].text = before + new_text[prefix:len(new_text)-suffix if suffix else None]
    for i in range(first + 1, last):
        paragraph.runs[i].text = ""
    if last != first:
        paragraph.runs[last].text = after
    else:
        paragraph.runs[first].text += after


class DocxProcessor(DocumentProcessor):
    def process(self, file_data, request):
        document = Document(io.BytesIO(file_data))
        for paragraph in _iter_paragraphs(document):
            old = "".join(run.text for run in paragraph.runs)
            if not old:
                continue
            lines = old.splitlines(keepends=True)
            new = "".join(convert_musical_text(line.rstrip("\r\n"), request) + line[len(line.rstrip("\r\n")):] for line in lines)
            _replace_runs(paragraph, new)
        output = io.BytesIO(); document.save(output)
        return output.getvalue()
