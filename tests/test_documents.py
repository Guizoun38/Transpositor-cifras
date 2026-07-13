import io

import pymupdf as fitz
import pytest
from docx import Document

from application.models import ConversionRequest
from application.use_cases import process_document
from domain.models import ConversionMode

REQUEST = ConversionRequest(ConversionMode.CHORDS_TO_CHORDS, "Ab", "C")


def _docx_text(data):
    doc = Document(io.BytesIO(data))
    return doc, "\n".join(p.text for p in doc.paragraphs)


def _pdf_words(document):
    return [word for page in document for word in page.get_text("words")]


def _assert_replacement_fits_before_bars(result, expected):
    converted = fitz.open(stream=result, filetype="pdf")
    words = _pdf_words(converted)
    replacement = next(word for word in words if word[4] == expected)
    bars = next(word for word in words if word[4] == "///")
    expected_width = fitz.get_text_length(expected, fontname="helv", fontsize=9)
    assert abs((replacement[2] - replacement[0]) - expected_width) < 0.1
    assert replacement[2] + 0.5 <= bars[0]
    assert abs((replacement[3] - replacement[1]) - (bars[3] - bars[1])) < 0.1
    assert bars[4] == "///"
    converted.close()


def test_generated_docx_integration():
    original = Document()
    original.add_paragraph("Pra. Giovanna - Tom Ab", style="Heading 1")
    original.add_paragraph("Ab Eb Fm Db")
    original.add_paragraph("Vem, Senhor")
    source = io.BytesIO(); original.save(source)
    result = process_document(source.getvalue(), ".docx", REQUEST)
    doc, text = _docx_text(result)
    assert "Pra. Giovanna - Tom C" in text
    assert "C G Am F" in text
    assert "Vem, Senhor" in text
    assert len(doc.paragraphs) == 3
    assert doc.paragraphs[0].style.name == "Heading 1"


@pytest.mark.parametrize(
    "conversion_request, intro, interlude, expected_intro, expected_interlude, expected_minister",
    [
        (
            ConversionRequest(ConversionMode.CHORDS_TO_CHORDS, "E", "C"),
            "| E /// |", "| A /// |", "| C /// |", "| F /// |", "Tom C",
        ),
        (
            ConversionRequest(ConversionMode.CHORDS_TO_NASHVILLE, "E", None),
            "| E /// |", "| A /// |", "| 1 /// |", "| 4 /// |", "Tom E",
        ),
        (
            ConversionRequest(ConversionMode.NASHVILLE_TO_CHORDS, None, "C"),
            "| 6 /// |", "| 4 /// |", "| A /// |", "| F /// |", "Tom C",
        ),
    ],
)
def test_docx_intro_interlude_and_key_headers(
    conversion_request, intro, interlude, expected_intro, expected_interlude, expected_minister
):
    original = Document()
    paragraph = original.add_paragraph()
    paragraph.add_run(
        "Artista/Banda: Fhop e Nívea Soares - Tom E; Bpm 72; 4/4\n"
        "Pra. Giovanna - Tom E\n"
        "[Intro]\n"
        f"{intro}\n"
        "[Interlúdio]\n"
        f"{interlude}"
    )
    source = io.BytesIO()
    original.save(source)

    result = process_document(source.getvalue(), ".docx", conversion_request)
    _, text = _docx_text(result)

    assert "Artista/Banda: Fhop e Nívea Soares - Tom E; Bpm 72; 4/4" in text
    assert f"Pra. Giovanna - {expected_minister}" in text
    assert expected_intro in text
    assert expected_interlude in text


def test_docx_chord_lines_with_performance_notes_convert_to_nashville():
    original = Document()
    original.add_paragraph(
        "B A F#m   (frase guitarra em A)\n"
        "C#m DM E A G#m E (A partir da 3ªx - Apenas Guitarra e baixo)\n"
        "| C#m /// | E /// | A/ G#m / | F#m / E / | (frase todos instrumentos Harmonia)\n"
        "(| DbM /// | B /// | E /// | E /// | - Apenas 2ªx Notas do Baixo)"
    )
    source = io.BytesIO()
    original.save(source)
    request = ConversionRequest(ConversionMode.CHORDS_TO_NASHVILLE, "E", None)

    result = process_document(source.getvalue(), ".docx", request)
    _, text = _docx_text(result)

    assert "5 4 2m   (frase guitarra em A)" in text
    assert "6m b7M 1 4 3m 1 (A partir da 3ªx - Apenas Guitarra e baixo)" in text
    assert "| 6m /// | 1 /// | 4/ 3m / | 2m / 1 / | (frase todos instrumentos Harmonia)" in text
    assert "(| 6M /// | 5 /// | 1 /// | 1 /// | - Apenas 2ªx Notas do Baixo)" in text


def test_docx_nashville_lines_with_performance_notes_convert_to_chords():
    original = Document()
    original.add_paragraph(
        "6m b7M 1 4 3m 1 (A partir da 3ªx - Apenas Guitarra e baixo)\n"
        "(| 6M /// | 5 /// | 1 /// | 1 /// | - Apenas 2ªx Notas do Baixo)\n"
        "| 6m /// | 1 /// | 4/ 3m / | 2m / 1 / | (frase todos instrumentos Harmonia)"
    )
    source = io.BytesIO()
    original.save(source)
    request = ConversionRequest(ConversionMode.NASHVILLE_TO_CHORDS, None, "C")

    result = process_document(source.getvalue(), ".docx", request)
    _, text = _docx_text(result)

    assert "Am BbM C F Em C (A partir da 3ªx - Apenas Guitarra e baixo)" in text
    assert "(| AM /// | G /// | C /// | C /// | - Apenas 2ªx Notas do Baixo)" in text
    assert "| Am /// | C /// | F/ Em / | Dm / C / | (frase todos instrumentos Harmonia)" in text


def test_generated_pdf_integration():
    source_doc = fitz.open(); page = source_doc.new_page()
    page.insert_text((72, 72), "Tom Ab\n\nAb Eb Fm Db\nVem, Senhor")
    source = source_doc.tobytes(); source_doc.close()
    result = process_document(source, ".pdf", REQUEST)
    converted = fitz.open(stream=result, filetype="pdf")
    words = [word[4] for word in _pdf_words(converted)]
    assert len(converted) == 1
    assert all(token in words for token in ("Tom", "C", "G", "Am", "F", "Vem,"))
    assert not all(token in words for token in ("Ab", "Eb", "Fm", "Db"))
    converted.close()


def test_generated_pdf_chords_to_nashville_remain_visible():
    source_doc = fitz.open(); page = source_doc.new_page()
    page.insert_text((72, 72), "Tom Ab")
    page.insert_text((72, 100), "Ab9 Ebsus4 Bbm7 Db")
    source = source_doc.tobytes(); source_doc.close()
    request = ConversionRequest(ConversionMode.CHORDS_TO_NASHVILLE, "Ab", None)
    result = process_document(source, ".pdf", request)
    converted = fitz.open(stream=result, filetype="pdf")
    words = [word[4] for word in _pdf_words(converted)]
    assert all(token in words for token in ("19", "5sus4", "2m7", "4"))
    assert not any(token in words for token in ("Ab9", "Ebsus4", "Bbm7", "Db"))
    converted.close()


def test_pdf_parenthesized_progression_with_inline_note_converts_to_nashville():
    source_doc = fitz.open(); page = source_doc.new_page()
    page.insert_text((72, 72), "(| DbM /// | B /// | E /// | E /// | - Apenas 2x Notas do Baixo)")
    source = source_doc.tobytes(); source_doc.close()
    request = ConversionRequest(ConversionMode.CHORDS_TO_NASHVILLE, "E", None)

    result = process_document(source, ".pdf", request)
    converted = fitz.open(stream=result, filetype="pdf")
    words = [word[4] for word in _pdf_words(converted)]

    assert all(token in words for token in ("6M", "5", "1", "Apenas", "Notas", "Baixo)"))
    assert not any(token in words for token in ("DbM", "B", "E"))
    converted.close()


def test_pdf_nashville_lines_with_performance_notes_convert_to_chords():
    source_doc = fitz.open(); page = source_doc.new_page()
    page.insert_text((72, 72), "6m b7M 1 4 3m 1 (A partir da 3x - Apenas Guitarra e baixo)")
    page.insert_text((72, 100), "(| 6M /// | 5 /// | 1 /// | 1 /// | - Apenas 2x Notas do Baixo)")
    page.insert_text((72, 128), "| 6m /// | 1 /// | 4/ 3m / | 2m / 1 / | (frase todos instrumentos Harmonia)")
    source = source_doc.tobytes(); source_doc.close()
    request = ConversionRequest(ConversionMode.NASHVILLE_TO_CHORDS, None, "C")

    result = process_document(source, ".pdf", request)
    converted = fitz.open(stream=result, filetype="pdf")
    words = [word[4] for word in _pdf_words(converted)]

    assert all(token in words for token in ("Am", "BbM", "AM", "G", "C", "F/", "Em", "Dm"))
    assert all(token in words for token in ("Apenas", "Guitarra", "2x", "Notas", "Harmonia)"))
    assert not any(token in words for token in ("6m", "b7M", "6M", "3m", "2m"))
    converted.close()


@pytest.mark.parametrize(
    "source_text, conversion_request, expected",
    [
        ("B ///", ConversionRequest(ConversionMode.CHORDS_TO_CHORDS, "E", "G#"), "D#"),
        ("B ///", ConversionRequest(ConversionMode.CHORDS_TO_CHORDS, "E", "Ab"), "Eb"),
        ("G ///", ConversionRequest(ConversionMode.CHORDS_TO_NASHVILLE, "Db", None), "b5"),
        ("5 ///", ConversionRequest(ConversionMode.NASHVILLE_TO_CHORDS, None, "G#"), "D#"),
        ("5 ///", ConversionRequest(ConversionMode.NASHVILLE_TO_CHORDS, None, "Ab"), "Eb"),
    ],
)
def test_pdf_wider_replacements_do_not_cover_measure_bars(source_text, conversion_request, expected):
    source_doc = fitz.open(); page = source_doc.new_page()
    page.insert_text((72, 72), source_text, fontsize=9)
    source = source_doc.tobytes(); source_doc.close()

    result = process_document(source, ".pdf", conversion_request)

    _assert_replacement_fits_before_bars(result, expected)


def test_pdf_untouched_lyrics_keep_original_formatting():
    source_doc = fitz.open(); page = source_doc.new_page()
    page.insert_text((72, 72), "Ab Eb Fm Db", fontsize=12, color=(1, 0.2, 0))
    page.insert_text((72, 100), "Vem, Senhor", fontsize=14, fontname="hebi", color=(0.2, 0.1, 0.7))
    source = source_doc.tobytes(); source_doc.close()

    original = fitz.open(stream=source, filetype="pdf")
    original_lyric = next(span for block in original[0].get_text("dict")["blocks"]
                          for line in block.get("lines", []) for span in line["spans"]
                          if span["text"] == "Vem, Senhor")
    result = process_document(source, ".pdf", REQUEST)
    converted = fitz.open(stream=result, filetype="pdf")
    converted_lyric = next(span for block in converted[0].get_text("dict")["blocks"]
                           for line in block.get("lines", []) for span in line["spans"]
                           if span["text"] == "Vem, Senhor")

    assert converted_lyric["font"] == original_lyric["font"]
    assert converted_lyric["size"] == original_lyric["size"]
    assert converted_lyric["color"] == original_lyric["color"]
    assert converted_lyric["bbox"] == original_lyric["bbox"]
    converted_words = [word[4] for word in converted[0].get_text("words")]
    assert all(token in converted_words for token in ("C", "G", "Am", "F"))
    original.close(); converted.close()


