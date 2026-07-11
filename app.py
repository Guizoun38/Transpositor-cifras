"""
Transpositor de Cifras em Word
--------------------------------
Aplicação Streamlit que importa um ficheiro .docx com uma cifra musical,
transpõe o tom indicado no cabeçalho e os acordes da música para um novo
tom, preservando toda a restante formatação e conteúdo do documento.
"""

import io
import re

import streamlit as st
from docx import Document


# ---------------------------------------------------------------------------
# Tabelas de notas
# ---------------------------------------------------------------------------

SHARP_NAMES = ["C", "C#", "D", "D#", "E", "F", "F#", "G", "G#", "A", "A#", "B"]
FLAT_NAMES = ["C", "Db", "D", "Eb", "E", "F", "Gb", "G", "Ab", "A", "Bb", "B"]

NOTE_TO_SEMITONE = {
    "C": 0, "B#": 0,
    "C#": 1, "Db": 1,
    "D": 2,
    "D#": 3, "Eb": 3,
    "E": 4, "Fb": 4,
    "F": 5, "E#": 5,
    "F#": 6, "Gb": 6,
    "G": 7,
    "G#": 8, "Ab": 8,
    "A": 9,
    "A#": 10, "Bb": 10,
    "B": 11, "Cb": 11,
}

# Tonalidades que tradicionalmente usam bemóis
FLAT_KEYS = {"F", "Bb", "Eb", "Ab", "Db", "Gb", "Cb",
             "Dm", "Gm", "Cm", "Fm", "Bbm", "Ebm", "Abm"}

TONES = [
    "C", "C#", "Db", "D", "D#", "Eb", "E", "F", "F#",
    "Gb", "G", "G#", "Ab", "A", "A#", "Bb", "B",
]


def uses_flats(key: str) -> bool:
    """Decide se a tonalidade de destino deve usar bemóis ou sustenidos."""
    root = key.rstrip("m")
    return root in FLAT_KEYS or key in FLAT_KEYS


def semitone_diff(original_key: str, new_key: str) -> int:
    root_o = original_key.rstrip("m")
    root_n = new_key.rstrip("m")
    if root_o not in NOTE_TO_SEMITONE or root_n not in NOTE_TO_SEMITONE:
        return 0
    return (NOTE_TO_SEMITONE[root_n] - NOTE_TO_SEMITONE[root_o]) % 12


def transpose_note_name(note: str, diff: int, flats: bool) -> str:
    if note not in NOTE_TO_SEMITONE:
        return note
    idx = (NOTE_TO_SEMITONE[note] + diff) % 12
    return FLAT_NAMES[idx] if flats else SHARP_NAMES[idx]


# ---------------------------------------------------------------------------
# Reconhecimento de acordes
# ---------------------------------------------------------------------------

# Grupo 1: nota principal (ex: C, C#, Db)
# Grupo 2: qualidade/extensão (ex: m7, sus4, maj7, add9, dim, aug, 9, 6...)
# Grupo 3: nota do baixo, se existir (acorde invertido, ex: C/E)
# Grupo 4: sufixo "m" do baixo, se existir (ex: C/Gm)
CHORD_RE = re.compile(
    r"^([A-G](?:#|b)?)"
    r"((?:maj|min|dim|aug|sus|add|m)?[0-9]*"
    r"(?:(?:maj|min|dim|aug|sus|add|m)[0-9]*)*[+\-]?)"
    r"(?:/([A-G](?:#|b)?)(m)?)?$"
)

# Tokens estruturais que podem aparecer numa linha de acordes sem serem
# acordes em si (barras de compasso, repetições, etc.)
STRUCTURAL_TOKENS = {"%", "-"}
STRUCTURAL_RE = re.compile(r"^[|/]+$")
REPEAT_RE = re.compile(r"^\d+x$", re.IGNORECASE)


def is_chord_token(token: str) -> bool:
    core = token.strip()
    if core == "" or core in STRUCTURAL_TOKENS:
        return True
    if STRUCTURAL_RE.fullmatch(core) or REPEAT_RE.fullmatch(core):
        return True
    core_bare = core.strip("|")
    return bool(CHORD_RE.fullmatch(core_bare))


def is_real_chord(token: str) -> bool:
    core = token.strip().strip("|")
    if core == "":
        return False
    return bool(CHORD_RE.fullmatch(core))


def is_chord_line(text: str) -> bool:
    """Deteção conservadora de linhas de acordes.

    Uma linha é considerada linha de acordes quando é composta
    maioritariamente por acordes reconhecidos, barras de compasso e
    símbolos de repetição. Linhas curtas com um único acorde (ex: "Fm")
    também são aceites. Linhas de letra que contenham por coincidência
    uma palavra igual a um acorde (ex: "A porta está aberta") são
    rejeitadas por exigirem pelo menos dois acordes reais quando a
    linha tem mais de dois tokens.
    """
    stripped = text.strip()
    if not stripped:
        return False
    tokens = stripped.split()
    if not tokens:
        return False

    total = len(tokens)
    chordish = sum(1 for tok in tokens if is_chord_token(tok))
    real_chords = sum(1 for tok in tokens if is_real_chord(tok))

    if real_chords == 0:
        return False

    if total <= 2:
        return (chordish / total) >= 0.5

    return real_chords >= 2 and (chordish / total) >= 0.35


def transpose_chord_token(token: str, diff: int, flats: bool) -> str:
    """Transpõe um único token de acorde, preservando extensão/baixo."""
    core = token
    prefix = ""
    suffix = ""

    # preserva barras verticais em volta do token (ex: "|Ab")
    while core.startswith("|"):
        prefix += "|"
        core = core[1:]
    while core.endswith("|"):
        suffix = "|" + suffix
        core = core[:-1]

    if core in STRUCTURAL_TOKENS or core == "":
        return token

    match = CHORD_RE.fullmatch(core)
    if not match:
        return token

    root, quality, bass, bass_suffix = match.groups()
    new_root = transpose_note_name(root, diff, flats)
    new_core = new_root + quality
    if bass:
        new_bass = transpose_note_name(bass, diff, flats)
        new_core += "/" + new_bass + (bass_suffix or "")

    return prefix + new_core + suffix


def chord_line_replacements(line: str, diff: int, flats: bool):
    """Devolve uma lista de (start, end, novo_texto) para os acordes
    reais encontrados numa linha, com posições relativas à linha."""
    replacements = []
    for m in re.finditer(r"\S+", line):
        token = m.group()
        if is_real_chord(token):
            new_token = transpose_chord_token(token, diff, flats)
            if new_token != token:
                replacements.append((m.start(), m.end(), new_token))
    return replacements


# ---------------------------------------------------------------------------
# Deteção / alteração da linha "Tom"
# ---------------------------------------------------------------------------

TOM_RE = re.compile(r"(Tom\s+)([A-G](?:#|b)?m?)\b")


def is_tom_line(text: str) -> bool:
    return TOM_RE.search(text) is not None


def _transpose_tom_note(note: str, diff: int, flats: bool) -> str:
    is_minor = note.endswith("m") and note not in NOTE_TO_SEMITONE
    root = note[:-1] if is_minor else note
    new_root = transpose_note_name(root, diff, flats)
    return new_root + ("m" if is_minor else "")


def tom_line_replacements(line: str, diff: int, flats: bool):
    """Devolve uma lista de (start, end, novo_texto) para a nota do tom
    encontrada na linha (apenas a nota é substituída, não o "Tom ")."""
    replacements = []
    for m in TOM_RE.finditer(line):
        note = m.group(2)
        new_note = _transpose_tom_note(note, diff, flats)
        if new_note != note:
            replacements.append((m.start(2), m.end(2), new_note))
    return replacements


# ---------------------------------------------------------------------------
# Aplicação da transposição a runs preservando formatação
# ---------------------------------------------------------------------------

def _classify_line(line: str) -> str:
    if is_tom_line(line):
        return "tom"
    if is_chord_line(line):
        return "chord"
    return "other"


def _paragraph_replacements(full_text: str, diff: int, flats: bool):
    """Calcula todas as substituições (start, end, novo_texto) a aplicar
    ao texto completo de um parágrafo, com posições globais.

    Um "parágrafo" do Word pode conter várias linhas visuais separadas
    por quebras de linha manuais (Shift+Enter), representadas pelo
    caráter "\\n" no texto concatenado dos runs. Por isso o texto é
    primeiro dividido em linhas para classificação e deteção dos
    acordes/tom a transpor.
    """
    lines = full_text.split("\n")
    replacements = []
    line_start = 0
    for line in lines:
        kind = _classify_line(line)
        if kind == "tom":
            for start, end, new_text in tom_line_replacements(line, diff, flats):
                replacements.append((line_start + start, line_start + end, new_text))
        elif kind == "chord":
            for start, end, new_text in chord_line_replacements(line, diff, flats):
                replacements.append((line_start + start, line_start + end, new_text))
        line_start += len(line) + 1  # +1 para o caráter "\n" removido pelo split
    return replacements


def _run_offsets(runs):
    offsets = []
    pos = 0
    for r in runs:
        offsets.append(pos)
        pos += len(r.text)
    return offsets


def _replace_span_in_runs(runs, start: int, end: int, new_text: str):
    """Substitui o intervalo [start, end) do texto concatenado dos runs
    por new_text, preservando a formatação: o texto novo é colocado no
    primeiro run afetado, e os runs seguintes afetados são ajustados ou
    esvaziados conforme necessário."""
    offsets = _run_offsets(runs)
    first_idx = None
    last_idx = None
    for i, off in enumerate(offsets):
        r_len = len(runs[i].text)
        r_start, r_end = off, off + r_len
        if r_end > start and r_start < end:
            if first_idx is None:
                first_idx = i
            last_idx = i

    if first_idx is None:
        return

    first_run = runs[first_idx]
    first_off = offsets[first_idx]
    before = first_run.text[: start - first_off]

    if last_idx == first_idx:
        after = first_run.text[end - first_off:]
        first_run.text = before + new_text + after
        return

    last_run = runs[last_idx]
    last_off = offsets[last_idx]
    after = last_run.text[end - last_off:]

    first_run.text = before + new_text
    last_run.text = after
    for i in range(first_idx + 1, last_idx):
        runs[i].text = ""


def process_paragraph(paragraph, diff: int, flats: bool):
    full_text = "".join(run.text for run in paragraph.runs)
    if not full_text.strip():
        return

    replacements = _paragraph_replacements(full_text, diff, flats)
    if not replacements:
        return

    # aplica de trás para a frente para não invalidar offsets anteriores
    replacements.sort(key=lambda r: r[0], reverse=True)
    for start, end, new_text in replacements:
        _replace_span_in_runs(paragraph.runs, start, end, new_text)


def iter_all_paragraphs(document: Document):
    """Devolve todos os parágrafos do documento: corpo, tabelas,
    cabeçalhos e rodapés."""
    paragraphs = list(document.paragraphs)

    def table_paragraphs(table):
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
                for t in cell.tables:
                    yield from table_paragraphs(t)

    for table in document.tables:
        paragraphs.extend(table_paragraphs(table))

    for section in document.sections:
        for part in (section.header, section.footer):
            paragraphs.extend(part.paragraphs)
            for table in part.tables:
                paragraphs.extend(table_paragraphs(table))

    return paragraphs


def transpose_document(file_bytes: bytes, original_key: str, new_key: str) -> bytes:
    document = Document(io.BytesIO(file_bytes))
    diff = semitone_diff(original_key, new_key)
    flats = uses_flats(new_key)

    for paragraph in iter_all_paragraphs(document):
        process_paragraph(paragraph, diff, flats)

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output.read()


# ---------------------------------------------------------------------------
# Interface Streamlit
# ---------------------------------------------------------------------------

st.set_page_config(page_title="Transpositor de Cifras", page_icon="🎸")
st.title("🎸 Transpositor de Cifras em Word")

uploaded_file = st.file_uploader("Carregar ficheiro Word (.docx)", type=["docx"])

col1, col2 = st.columns(2)
with col1:
    original_key = st.selectbox("Tom original", TONES, index=TONES.index("Ab"))
with col2:
    new_key = st.selectbox("Novo tom", TONES, index=TONES.index("C"))

if "output_bytes" not in st.session_state:
    st.session_state.output_bytes = None
    st.session_state.output_name = None

if st.button("Transpor ficheiro", type="primary", disabled=uploaded_file is None):
    if uploaded_file is not None:
        file_bytes = uploaded_file.getvalue()
        try:
            output_bytes = transpose_document(file_bytes, original_key, new_key)
            base_name = uploaded_file.name.rsplit(".", 1)[0]
            output_name = f"{base_name} - Tom {new_key}.docx"
            st.session_state.output_bytes = output_bytes
            st.session_state.output_name = output_name
            st.success("Ficheiro transposto com sucesso!")
        except Exception as exc:  # noqa: BLE001
            st.error(f"Ocorreu um erro ao processar o ficheiro: {exc}")

if st.session_state.output_bytes:
    st.download_button(
        "Descarregar ficheiro",
        data=st.session_state.output_bytes,
        file_name=st.session_state.output_name,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
